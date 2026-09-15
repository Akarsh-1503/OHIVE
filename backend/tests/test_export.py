"""Excel and CSV export: open the real artefacts and assert what a reviewer would see."""

from __future__ import annotations

import csv
import io
import re
from datetime import UTC, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.export import (
    COLUMNS,
    FIRST_DATA_ROW,
    HEADER_ROW,
    ROW_DUPLICATE,
    ROW_FAILED,
    ROW_NEEDS_REVIEW,
    build_csv,
    build_workbook,
    export_filename,
    short_id,
)
from app.models import Batch, Confidence, Lead
from tests.conftest import upload_parts
from tests.test_contract import wait_for_batch

BATCH_ID = "3c2a1b9f4e6d47a1b2c3d4e5f6a7b8c9"


def lead(card_id: str, **over: object) -> Lead:
    base: dict[str, object] = {
        "card_id": card_id,
        "batch_id": BATCH_ID,
        "filename": f"{card_id}.jpg",
        "status": "completed",
        "first_name": "Priya",
        "last_name": "Raghavan",
        "job_title": "VP of Partnerships",
        "company": "Northwind Robotics",
        "location": "Bengaluru, KA, India",
        "phone": "+91 80 4718 2200",
        "phone_e164": "+918047182200",
        "email": "priya@northwind.io",
        "website": "northwind.io",
        "confidence": Confidence(
            first_name=0.97, last_name=0.96, job_title=0.91, company=0.98,
            location=0.74, phone=0.99, email=0.99, website=0.88,
        ),
        "overall_confidence": 0.93,
        "raw_text": "Priya Raghavan\nVP of Partnerships\nNorthwind Robotics",
        "processing_ms": 2140,
    }
    return Lead(**(base | over))


@pytest.fixture
def batch() -> Batch:
    return Batch(
        batch_id=BATCH_ID,
        status="partial",
        total=4,
        completed=3,
        failed=1,
        pending=0,
        created_at="2026-09-14T09:12:03Z",
        finished_at="2026-09-14T09:12:21Z",
        elapsed_ms=18422,
        leads=[
            lead("card1"),
            lead("card2", duplicate_of="card1", processing_ms=1980),
            lead(
                "card3",
                status="needs_review",
                overall_confidence=0.41,
                quality_flags=["blurry", "no_email"],
                email=None,
                website=None,
                processing_ms=3100,
            ),
            lead(
                "card4",
                status="failed",
                error="vlm unavailable after 3 attempts",
                first_name=None, last_name=None, job_title=None, company=None,
                location=None, phone=None, phone_e164=None, email=None, website=None,
                confidence=Confidence(),
                overall_confidence=0.0,
                raw_text=None,
                processing_ms=310,
            ),
        ],
    )


@pytest.fixture
def raw() -> dict[str, dict[str, object]]:
    return {
        "card1": {"raw_text": "Priya Raghavan", "fields": {"company": {"value": "Northwind"}}},
    }


MODEL = "Qwen/Qwen2.5-VL-3B-Instruct"


def workbook(batch: Batch, raw: dict[str, dict[str, object]], **kwargs: object) -> Workbook:
    return load_workbook(io.BytesIO(build_workbook(batch, raw, "stub", MODEL, **kwargs)))


def summary_values(sheet: Worksheet) -> dict[str, object]:
    return {
        row[0]: row[1]
        for row in sheet.iter_rows(min_row=HEADER_ROW, max_col=2, values_only=True)
        if isinstance(row[0], str)
    }


# --------------------------------------------------------------------------- workbook


def test_workbook_has_the_three_sheets(batch: Batch, raw) -> None:
    assert workbook(batch, raw).sheetnames == ["Leads", "Summary", "Raw"]


def test_branded_title_band(batch: Batch, raw) -> None:
    sheet = workbook(batch, raw)["Leads"]
    assert "LeadForge" in sheet["A1"].value
    caption = sheet["A2"].value
    assert short_id(BATCH_ID) in caption
    assert "4 source files" in caption
    assert "Qwen/Qwen2.5-VL-3B-Instruct via stub" in caption
    assert "generated" in caption


def test_header_row_is_frozen_filtered_and_complete(batch: Batch, raw) -> None:
    sheet = workbook(batch, raw)["Leads"]
    headers = [cell.value for cell in sheet[HEADER_ROW]]
    assert headers == [title for title, _ in COLUMNS]
    assert sheet.freeze_panes == f"A{FIRST_DATA_ROW}"
    assert sheet.auto_filter.ref == f"A{HEADER_ROW}:P{FIRST_DATA_ROW + 3}"
    assert sheet.print_area
    assert all(sheet.column_dimensions[chr(ord("A") + i)].width >= 9 for i in range(len(COLUMNS)))


def test_rows_carry_the_lead_data(batch: Batch, raw) -> None:
    sheet = workbook(batch, raw)["Leads"]
    row = {title: sheet.cell(row=FIRST_DATA_ROW, column=index).value
           for index, (title, _) in enumerate(COLUMNS, 1)}
    assert row["#"] == 1
    assert row["Status"] == "completed"
    assert row["First Name"] == "Priya"
    assert row["Company"] == "Northwind Robotics"
    assert row["Phone (E.164)"] == "+918047182200"
    assert row["Confidence"] == 0.93
    assert row["Source File"] == "card1.jpg"


def test_hyperlinks_are_clickable(batch: Batch, raw) -> None:
    sheet = workbook(batch, raw)["Leads"]
    assert sheet.cell(row=FIRST_DATA_ROW, column=10).hyperlink.target == "mailto:priya@northwind.io"
    assert sheet.cell(row=FIRST_DATA_ROW, column=9).hyperlink.target == "tel:+918047182200"
    assert sheet.cell(row=FIRST_DATA_ROW, column=11).hyperlink.target == "https://northwind.io"


def test_e164_is_stored_as_text_so_excel_keeps_the_plus(batch: Batch, raw) -> None:
    sheet = workbook(batch, raw)["Leads"]
    cell = sheet.cell(row=FIRST_DATA_ROW, column=9)
    assert cell.number_format == "@"
    assert isinstance(cell.value, str) and cell.value.startswith("+")


def test_confidence_column_has_a_colour_scale(batch: Batch, raw) -> None:
    sheet = workbook(batch, raw)["Leads"]
    ranges = {str(rng.sqref): rng.rules for rng in sheet.conditional_formatting}
    target = f"L{FIRST_DATA_ROW}:L{FIRST_DATA_ROW + 3}"
    assert target in ranges
    assert ranges[target][0].type == "colorScale"
    assert len(ranges[target][0].colorScale.cfvo) == 3


def test_rows_are_tinted_by_review_state(batch: Batch, raw) -> None:
    sheet = workbook(batch, raw)["Leads"]

    def fill(offset: int) -> str:
        return sheet.cell(row=FIRST_DATA_ROW + offset, column=3).fill.fgColor.rgb

    assert fill(0) in ("00000000", None)  # clean row is untinted
    assert fill(1).endswith(ROW_DUPLICATE)
    assert fill(2).endswith(ROW_NEEDS_REVIEW)
    assert fill(3).endswith(ROW_FAILED)


def test_duplicate_column_points_at_the_anchor(batch: Batch, raw) -> None:
    sheet = workbook(batch, raw)["Leads"]
    assert sheet.cell(row=FIRST_DATA_ROW + 1, column=14).value == short_id("card1")
    # openpyxl writes an empty string as a blank cell, which reads back as None.
    assert sheet.cell(row=FIRST_DATA_ROW, column=14).value is None


# --------------------------------------------------------------------------- summary


def test_summary_numbers(batch: Batch, raw) -> None:
    values = summary_values(workbook(batch, raw)["Summary"])
    assert values["Batch ID"] == BATCH_ID
    assert values["Batch status"] == "partial"
    assert values["Source files"] == 4
    assert values["Rows exported"] == 4
    assert values["Completed"] == 2
    assert values["Needs review"] == 1
    assert values["Failed"] == 1
    assert values["Duplicates found"] == 1
    assert values["Mean confidence"] == pytest.approx((0.93 + 0.93 + 0.41) / 3, abs=1e-4)
    assert values["Batch elapsed (s)"] == pytest.approx(18.42, abs=0.01)
    assert values["Total processing time (s)"] == pytest.approx(7.53, abs=0.01)
    assert values["VLM provider"] == "stub"
    assert values["VLM model"] == "Qwen/Qwen2.5-VL-3B-Instruct"


def test_summary_field_fill_rates(batch: Batch, raw) -> None:
    sheet = workbook(batch, raw)["Summary"]
    rows = {
        row[0]: (row[1], row[2], row[3])
        for row in sheet.iter_rows(min_row=HEADER_ROW, max_col=4, values_only=True)
        if isinstance(row[0], str) and row[0].islower()
    }
    assert set(rows) == {
        "first_name", "last_name", "job_title", "company",
        "location", "phone", "email", "website",
    }
    # Three scored cards; two of them carry an email.
    assert rows["email"] == (2, 3, pytest.approx(2 / 3, abs=1e-4))
    assert rows["company"] == (3, 3, 1.0)


# --------------------------------------------------------------------------- raw sheet


def test_raw_sheet_carries_the_transcription_and_json(batch: Batch, raw) -> None:
    sheet = workbook(batch, raw)["Raw"]
    assert [cell.value for cell in sheet[HEADER_ROW]] == [
        "Card ID", "Source File", "Status", "Confidence", "Raw Text", "Raw JSON"
    ]
    assert sheet.cell(row=FIRST_DATA_ROW, column=1).value == "card1"
    assert "Priya Raghavan" in sheet.cell(row=FIRST_DATA_ROW, column=5).value
    assert "Northwind" in sheet.cell(row=FIRST_DATA_ROW, column=6).value
    assert sheet.cell(row=FIRST_DATA_ROW + 1, column=6).value is None  # no raw json stored


# --------------------------------------------------------------------------- filtering & csv


def test_include_low_confidence_false_keeps_only_clean_rows(batch: Batch, raw) -> None:
    sheet = workbook(batch, raw, include_low_confidence=False)["Leads"]
    statuses = [
        sheet.cell(row=FIRST_DATA_ROW + offset, column=2).value for offset in range(3)
    ]
    assert statuses == ["completed", "completed", None]


def test_include_duplicates_false_hides_the_duplicate_row(batch: Batch, raw) -> None:
    sheet = workbook(batch, raw, include_duplicates=False)["Leads"]
    exported = [
        sheet.cell(row=FIRST_DATA_ROW + offset, column=15).value for offset in range(4)
    ]
    # card2 is the duplicate of card1 and is the only row dropped.
    assert exported == ["card1.jpg", "card3.jpg", "card4.jpg", None]
    assert all(
        sheet.cell(row=FIRST_DATA_ROW + offset, column=14).value is None for offset in range(3)
    )


def test_summary_still_counts_duplicates_when_they_are_hidden(batch: Batch, raw) -> None:
    """The point of the flag is a clean import file, not a flattering report."""
    values = summary_values(workbook(batch, raw, include_duplicates=False)["Summary"])
    assert values["Duplicates found"] == 1
    assert values["Duplicates hidden from Leads sheet"] == 1
    assert values["Rows exported"] == 3
    assert values["Source files"] == 4
    # Every other count is still over the whole batch.
    assert values["Completed"] == 2
    assert values["Needs review"] == 1
    assert values["Failed"] == 1
    assert values["Mean confidence"] == pytest.approx((0.93 + 0.93 + 0.41) / 3, abs=1e-4)


def test_summary_reports_no_hidden_rows_by_default(batch: Batch, raw) -> None:
    values = summary_values(workbook(batch, raw)["Summary"])
    assert values["Duplicates hidden from Leads sheet"] == 0
    assert values["Rows exported"] == 4


def test_both_export_filters_compose(batch: Batch, raw) -> None:
    sheet = workbook(
        batch, raw, include_low_confidence=False, include_duplicates=False
    )["Leads"]
    exported = [
        sheet.cell(row=FIRST_DATA_ROW + offset, column=15).value for offset in range(2)
    ]
    assert exported == ["card1.jpg", None]


def test_export_filename_pattern() -> None:
    name = export_filename(BATCH_ID, "xlsx")
    stamp = datetime.now(UTC).strftime("%Y%m%d")
    assert name == f"leadforge-3c2a1b9f-{stamp}.xlsx"
    assert re.fullmatch(r"leadforge-[0-9a-f]{8}-\d{8}\.csv", export_filename(BATCH_ID, "csv"))


def test_csv_export(batch: Batch) -> None:
    payload = build_csv(batch)
    assert payload.startswith(b"\xef\xbb\xbf")  # BOM keeps Excel happy with UTF-8

    rows = list(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))
    assert len(rows) == 4
    assert rows[0]["email"] == "priya@northwind.io"
    assert rows[0]["phone_e164"] == "+918047182200"
    assert rows[0]["confidence_company"] == "0.98"
    assert rows[1]["duplicate_of"] == "card1"
    assert rows[2]["quality_flags"] == "blurry, no_email"
    assert rows[3]["error"] == "vlm unavailable after 3 attempts"
    assert "confidence" not in rows[0]  # flattened into confidence_<field>


def csv_rows(payload: bytes) -> list[dict[str, str]]:
    return list(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))))


def test_csv_filtering(batch: Batch) -> None:
    rows = csv_rows(build_csv(batch, include_low_confidence=False))
    assert [row["card_id"] for row in rows] == ["card1", "card2"]


def test_csv_include_duplicates_false(batch: Batch) -> None:
    rows = csv_rows(build_csv(batch, include_duplicates=False))
    assert [row["card_id"] for row in rows] == ["card1", "card3", "card4"]
    assert all(row["duplicate_of"] == "" for row in rows)


def test_csv_both_filters_compose(batch: Batch) -> None:
    rows = csv_rows(build_csv(batch, include_low_confidence=False, include_duplicates=False))
    assert [row["card_id"] for row in rows] == ["card1"]


def test_csv_leaves_processing_ms_empty_for_a_card_that_never_ran(batch: Batch) -> None:
    queued = lead("card5", status="queued", processing_ms=None, raw_text=None)
    rows = csv_rows(build_csv(batch.model_copy(update={"leads": [*batch.leads, queued]})))
    assert rows[-1]["processing_ms"] == ""


# --------------------------------------------------------------------------- over HTTP


def test_export_query_params_are_honoured_over_http(client: TestClient) -> None:
    """Regression: the frontend sends `include_duplicates`, so it must not be ignored."""
    batch_id = client.post(
        "/api/v1/batches", files=upload_parts(["scan_dup_a.jpg", "scan_dup_b.jpg", "solo_01.jpg"])
    ).json()["batch_id"]
    batch = wait_for_batch(client, batch_id)
    duplicates = [lead for lead in batch["leads"] if lead["duplicate_of"]]
    assert len(duplicates) == 1, "fixture should produce exactly one duplicate"

    everything = load_workbook(
        io.BytesIO(client.get(f"/api/v1/batches/{batch_id}/export.xlsx").content)
    )
    filtered = load_workbook(
        io.BytesIO(
            client.get(
                f"/api/v1/batches/{batch_id}/export.xlsx",
                params={"include_duplicates": "false"},
            ).content
        )
    )
    assert everything["Leads"].max_row - filtered["Leads"].max_row == 1
    assert summary_values(filtered["Summary"])["Duplicates found"] == 1
    assert summary_values(filtered["Summary"])["Rows exported"] == 2

    csv_all = csv_rows(client.get(f"/api/v1/batches/{batch_id}/export.csv").content)
    csv_filtered = csv_rows(
        client.get(
            f"/api/v1/batches/{batch_id}/export.csv", params={"include_duplicates": "false"}
        ).content
    )
    assert len(csv_all) == 3
    assert len(csv_filtered) == 2


def test_export_endpoints_serve_downloadable_files(client: TestClient) -> None:
    batch_id = client.post(
        "/api/v1/batches", files=upload_parts(["card_01.jpg", "card_02.jpg"])
    ).json()["batch_id"]
    wait_for_batch(client, batch_id)

    xlsx = client.get(f"/api/v1/batches/{batch_id}/export.xlsx")
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert f'filename="leadforge-{short_id(batch_id)}-' in xlsx.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(xlsx.content))["Leads"]
    assert sheet.cell(row=FIRST_DATA_ROW, column=15).value == "card_01.jpg"

    csv_response = client.get(f"/api/v1/batches/{batch_id}/export.csv")
    assert csv_response.status_code == 200
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert len(list(csv.DictReader(io.StringIO(csv_response.text.lstrip("﻿"))))) == 2


def test_export_of_an_unknown_batch_is_404(client: TestClient) -> None:
    response = client.get("/api/v1/batches/nope/export.xlsx")
    assert response.status_code == 404
    assert response.json()["code"] == "BATCH_NOT_FOUND"
